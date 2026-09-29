"""
services/prep_export_service.py

Mentify Prep Export Service:
Generates publication-ready PDF and DOCX documents with robust formatting for:
- Mathematical equations (LaTeX & Unicode symbol formatting)
- Markdown tables (rendered as native ReportLab and Word tables with teal headers)
- Fenced code blocks (R / Python in dark terminal styling and Courier/Consolas fonts)
- Theorems, definitions, authentic exam questions, and verified step-by-step proofs.
"""
from __future__ import annotations

import io
import json
import logging
import os
import re
import tempfile
from datetime import datetime

from django.conf import settings
logger = logging.getLogger(__name__)

try:
    import docx
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
except ImportError:
    docx = None
    Document = None
    Inches = Pt = RGBColor = None
    WD_ALIGN_PARAGRAPH = None
    parse_xml = nsdecls = None

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
        HRFlowable,
        KeepTogether,
    )
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
except ImportError:
    colors = None
    A4 = None
    getSampleStyleSheet = ParagraphStyle = None
    SimpleDocTemplate = Paragraph = Spacer = Table = TableStyle = HRFlowable = KeepTogether = None
    pdfmetrics = TTFont = None


# ─── Font Registration for ReportLab (Unicode & Math Support) ─────────────────

def _register_system_fonts():
    """Register TrueType fonts on Windows for proper Unicode math & code rendering."""
    try:
        font_map = {
            "Arial": r"C:\Windows\Fonts\arial.ttf",
            "Arial-Bold": r"C:\Windows\Fonts\arialbd.ttf",
            "Arial-Italic": r"C:\Windows\Fonts\ariali.ttf",
            "Consolas": r"C:\Windows\Fonts\consola.ttf",
            "Consolas-Bold": r"C:\Windows\Fonts\consolab.ttf",
        }
        for name, path in font_map.items():
            if os.path.exists(path) and name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, path))
    except Exception:
        pass

_register_system_fonts()

FONT_BODY = "Arial" if "Arial" in pdfmetrics.getRegisteredFontNames() else "Helvetica"
FONT_BOLD = "Arial-Bold" if "Arial-Bold" in pdfmetrics.getRegisteredFontNames() else "Helvetica-Bold"
FONT_ITALIC = "Arial-Italic" if "Arial-Italic" in pdfmetrics.getRegisteredFontNames() else "Helvetica-Oblique"
FONT_CODE = "Consolas" if "Consolas" in pdfmetrics.getRegisteredFontNames() else "Courier"


# ─── Mathematical Formula Beautifier & Sanitizer ─────────────────────────────

def beautify_math_formula(latex_str: str) -> str:
    """
    Transforms LaTeX mathematical expressions into publication-grade, readable Unicode math.
    Handles greek letters, operators, fractions, matrices, superscripts, and subscripts.
    """
    if not latex_str:
        return ""
    s = latex_str.strip()
    if s.startswith("$$") and s.endswith("$$") and len(s) > 2:
        s = s[2:-2].strip()
    elif s.startswith("$") and s.endswith("$") and len(s) > 1:
        s = s[1:-1].strip()
    elif s.startswith("\\[") and s.endswith("\\]"):
        s = s[2:-2].strip()
    elif s.startswith("\\(") and s.endswith("\\)"):
        s = s[2:-2].strip()

    # Greek letters
    greek = {
        r"\alpha": "α", r"\beta": "β", r"\gamma": "γ", r"\delta": "δ",
        r"\epsilon": "ε", r"\varepsilon": "ε", r"\zeta": "ζ", r"\eta": "η",
        r"\theta": "θ", r"\iota": "ι", r"\kappa": "κ", r"\lambda": "λ",
        r"\mu": "μ", r"\nu": "ν", r"\xi": "ξ", r"\pi": "π",
        r"\rho": "ρ", r"\sigma": "σ", r"\tau": "τ", r"\phi": "φ",
        r"\chi": "χ", r"\psi": "ψ", r"\omega": "ω",
        r"\Gamma": "Γ", r"\Delta": "Δ", r"\Theta": "Θ", r"\Lambda": "Λ",
        r"\Sigma": "Σ", r"\Phi": "Φ", r"\Psi": "Ψ", r"\Omega": "Ω",
    }
    for k, v in greek.items():
        s = re.sub(re.escape(k) + r"(?![a-zA-Z])", v, s)

    # Blackboard bold & Mathcal
    bb_map = {
        r"\mathbb{R}": "ℝ", r"\mathbb{P}": "ℙ", r"\mathbb{E}": "𝔼",
        r"\mathbb{N}": "ℕ", r"\mathbb{Z}": "ℤ", r"\mathbb{Q}": "ℚ",
        r"\mathbb{C}": "ℂ", r"\mathcal{F}": "ℱ", r"\mathcal{B}": "ℬ",
    }
    for k, v in bb_map.items():
        s = s.replace(k, v)

    # Mathematical Operators & Relations
    ops = {
        r"\le": "≤", r"\leq": "≤", r"\ge": "≥", r"\geq": "≥",
        r"\neq": "≠", r"\ne": "≠", r"\times": "×", r"\pm": "±",
        r"\cdot": "·", r"\dots": "…", r"\cdots": "…", r"\ddots": "⋱",
        r"\infty": "∞", r"\subset": "⊂", r"\subseteq": "⊆",
        r"\cap": "∩", r"\cup": "∪", r"\in": "∈", r"\notin": "∉",
        r"\forall": "∀", r"\exists": "∃", r"\rightarrow": "→",
        r"\to": "→", r"\implies": " ⟹ ", r"\iff": " ⟺ ",
        r"\blacksquare": "■", r"\square": "□", r"\approx": "≈",
        r"\sim": "~", r"\equiv": "≡", r"\partial": "∂",
        r"\top": "ᵀ", r"\prime": "′",
    }
    for k, v in ops.items():
        s = re.sub(re.escape(k) + r"(?![a-zA-Z])", v, s)

    # Text and operator formatting
    s = re.sub(r"\\text\{([^{}]+)\}", r" \1 ", s)
    s = re.sub(r"\\(?:operatorname|mathrm|mathbf|boldsymbol|mathit)\{([^{}]+)\}", r"\1", s)
    s = s.replace(r"^\top", "ᵀ").replace(r"^\intercal", "ᵀ").replace(r"\top", "ᵀ")
    s = s.replace(r"\otimes", " ⊗ ").replace(r"\odot", " ⊙ ")

    # Common mathematical functions & operators
    funcs = [
        "det", "ln", "log", "exp", "sin", "cos", "tan", "dim", "ker",
        "deg", "max", "min", "sup", "inf", "lim", "var", "cov", "corr",
        "rank", "diag", "trace", "tr", "span", "nullity",
    ]
    for f in funcs:
        s = re.sub(r"\\" + f + r"(?![a-zA-Z])", f, s)

    # Sum, product, integral, square roots
    s = s.replace(r"\sum", "∑").replace(r"\prod", "∏").replace(r"\int", "∫")
    s = re.sub(r"\\sqrt\{([^{}]+)\}", r"√(\1)", s)
    s = re.sub(r"\\sqrt\[(\d+)\]\{([^{}]+)\}", r"^\1√(\2)", s)

    # Fractions: \frac{a}{b} -> (a / b)
    s = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1 / \2)", s)
    s = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1 / \2)", s)

    # Superscripts & Subscripts Unicode
    sups = {'0': '⁰', '1': '¹', '2': '²', '3': '³', '4': '⁴', '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹', '+': '⁺', '-': '⁻', 'n': 'ⁿ', 'i': 'ⁱ', 'k': 'ᵏ', 't': 'ᵗ'}
    subs = {'0': '₀', '1': '₁', '2': '₂', '3': '₃', '4': '₄', '5': '₅', '6': '₆', '7': '₇', '8': '₈', '9': '₉', '+': '₊', '-': '₋', 'n': 'ₙ', 'i': 'ᵢ', 'j': 'ⱼ', 'k': 'ₖ', 'x': 'ₓ'}
    for k, v in sups.items():
        s = s.replace(f"^{k}", v).replace(f"^{{{k}}}", v)
    for k, v in subs.items():
        s = s.replace(f"_{k}", v).replace(f"_{{{k}}}", v)

    # Matrix cleaning: \begin{pmatrix} a & b \\ c & d \end{pmatrix} -> [ a  b ;  c  d ]
    s = re.sub(
        r"\\begin\{(?:p|b|v|B|V)?matrix\}([\s\S]*?)\\end\{(?:p|b|v|B|V)?matrix\}",
        lambda m: "[ " + re.sub(r"\\\\|\\cr", " ; ", m.group(1)).replace("&", "  ").strip() + " ]",
        s,
    )
    s = s.replace(r"\,", " ").replace(r"\;", " ").replace(r"\quad", "  ").replace(r"\qquad", "   ")
    return s.strip()


def _format_markdown_for_reportlab(text: str) -> str:
    """Format markdown text with inline math and code tags for ReportLab Paragraphs."""
    if not text:
        return ""

    # 1. Escape XML characters in raw text FIRST
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    # 2. Extract inline code `...` and replace with placeholders so code like `$values` is not matched as math
    code_tokens = []
    def _save_code(m):
        idx = len(code_tokens)
        code_tokens.append(f'<font face="{FONT_CODE}" color="#0f766e">{m.group(1)}</font>')
        return f"%%%CODE_TOKEN_{idx}%%%"

    text = re.sub(r"`([^`\n\r]+?)`", _save_code, text)

    # 3. Bold markdown **text**
    text = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", text)
    # 4. Italic markdown *text*
    text = re.sub(r"\*(.*?)\*", r"<i>\1</i>", text)

    # 5. Replace inline math $...$ with beautified Unicode math
    def _math_sub(match):
        inner = match.group(1)
        if "%%%CODE_TOKEN_" in inner:
            return match.group(0)
        b_clean = beautify_math_formula(inner)
        b_clean = b_clean.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return f"<b><i>{b_clean}</i></b>"

    text = re.sub(r"\$([^\$\n\r]+?)\$", _math_sub, text)

    # 6. Restore code tokens
    for idx, token in enumerate(code_tokens):
        text = text.replace(f"%%%CODE_TOKEN_{idx}%%%", token)

    return text


# ─── Markdown Document Block Parser ──────────────────────────────────────────

def _parse_markdown_into_blocks(raw_text: str) -> list[dict]:
    """
    Parses markdown content into structured blocks:
    - heading (level, text)
    - code (lang, lines)
    - table (headers, rows)
    - math (formula)
    - blockquote (text)
    - paragraph (text)
    """
    lines = raw_text.replace("\r\n", "\n").split("\n")
    blocks = []
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # 1. Fenced Code Block
        if stripped.startswith("```"):
            lang = stripped.lstrip("`").strip() or "R"
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < n and lines[i].strip().startswith("```"):
                i += 1
            blocks.append({"type": "code", "lang": lang, "lines": code_lines})
            continue

        # 2. Display Math Block ($$...$$ or \[...\])
        if stripped.startswith("$$") or stripped.startswith("\\["):
            if (stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) > 2) or \
               (stripped.startswith("\\[") and stripped.endswith("\\]")):
                formula = stripped
                i += 1
            else:
                math_lines = [stripped.lstrip("$").lstrip("\\[")]
                i += 1
                while i < n and not (lines[i].strip().endswith("$$") or lines[i].strip().endswith("\\]")):
                    math_lines.append(lines[i].strip())
                    i += 1
                if i < n:
                    math_lines.append(lines[i].strip().rstrip("$").rstrip("\\]"))
                    i += 1
                formula = " ".join(math_lines)
            blocks.append({"type": "math", "formula": formula})
            continue

        # 3. Markdown Table (lines starting and ending with |)
        if stripped.startswith("|") and stripped.endswith("|") and "|" in stripped[1:-1]:
            table_lines = []
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            rows = []
            for tl in table_lines:
                # Skip separator line like |---|---|
                if re.match(r"^\|[\s\-:|]+\|$", tl):
                    continue
                cells = [c.strip() for c in tl.strip("|").split("|")]
                rows.append(cells)

            if rows:
                headers = rows[0]
                data_rows = rows[1:] if len(rows) > 1 else []
                blocks.append({"type": "table", "headers": headers, "rows": data_rows})
            continue

        # 4. Headings
        if stripped.startswith("#"):
            match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if match:
                level = len(match.group(1))
                text = match.group(2).strip()
                blocks.append({"type": "heading", "level": level, "text": text})
                i += 1
                continue

        # 5. Blockquote (Theorems / Definitions)
        if stripped.startswith(">"):
            bq_lines = []
            while i < n and lines[i].strip().startswith(">"):
                bq_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append({"type": "blockquote", "text": " ".join(bq_lines)})
            continue

        # 6. Standard Paragraph or Bullet Point
        blocks.append({"type": "paragraph", "text": stripped})
        i += 1

    return blocks


# ─── 1. High-Fidelity PDF Export Engine (Playwright + Chrome) ─────────────────

def _render_html_with_playwright(
    html_content: str,
    header_text: str = "Mentify Academic Prep",
    timeout_ms: int = 15000,
) -> bytes:
    """
    Renders an HTML string to a publication-ready vector PDF using headless Chrome/Edge and Playwright.
    Returns empty bytes b"" if Playwright or Chrome is unavailable, falling back seamlessly.
    """
    chrome_candidates = [
        os.environ.get("CHROME_PATH", ""),
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    ]
    chrome_path = next((p for p in chrome_candidates if p and os.path.exists(p)), None)
    if not chrome_path:
        return b""

    tmp_path = None
    try:
        from playwright.sync_api import sync_playwright

        with tempfile.NamedTemporaryFile(suffix=".html", delete=False, mode="w", encoding="utf-8") as f:
            f.write(html_content)
            tmp_path = f.name

        norm_path = tmp_path.replace("\\", "/")
        file_url = f"file:///{norm_path}"

        with sync_playwright() as p:
            browser = p.chromium.launch(
                executable_path=chrome_path,
                args=["--no-sandbox", "--disable-gpu", "--disable-dev-shm-usage"],
            )
            page = browser.new_page()
            page.goto(file_url, wait_until="networkidle", timeout=timeout_ms)
            try:
                page.wait_for_selector(".render-complete", state="attached", timeout=8000)
            except Exception:
                pass

            safe_header = header_text.replace("<", "&lt;").replace(">", "&gt;")
            pdf_bytes = page.pdf(
                format="A4",
                margin={"top": "18mm", "bottom": "18mm", "left": "16mm", "right": "16mm"},
                print_background=True,
                display_header_footer=True,
                header_template=(
                    f'<div style="font-size: 8pt; font-family: -apple-system, BlinkMacSystemFont, sans-serif; '
                    f'color: #94a3b8; width: 100%; text-align: right; padding-right: 16mm;">'
                    f'{safe_header} &bull; Mentify Academic Prep</div>'
                ),
                footer_template=(
                    '<div style="font-size: 8pt; font-family: -apple-system, BlinkMacSystemFont, sans-serif; '
                    'color: #94a3b8; width: 100%; text-align: center;">'
                    'Page <span class="pageNumber"></span> of <span class="totalPages"></span></div>'
                ),
            )
            browser.close()
            return pdf_bytes
    except Exception as exc:
        logger.warning("Playwright PDF generation failed, falling back to ReportLab: %s", exc)
        return b""
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass


def _build_topic_notes_html(
    course_code: str,
    topic_title: str,
    notes_content: str,
    authentic_questions: list | None = None,
    practice_questions: list | None = None,
    level: str = "level_2",
) -> str:
    """Builds complete standalone HTML with KaTeX, marked, Prism.js, and print styling for topic notes."""
    from services.prep_ai_router import normalize_math_delimiters

    full_markdown = notes_content.strip()

    if authentic_questions:
        full_markdown += "\n\n---\n\n## Authentic Examination Problems & Verified Proofs\n\n"
        for q in authentic_questions:
            num = q.get("number", 1)
            marks = q.get("marks", 10)
            paper = q.get("paper_title", "University Examination")
            year = q.get("year", "")
            q_latex = q.get("question_latex", "")
            sol_latex = q.get("solution_latex", "")
            full_markdown += f"### Question {num} [{marks} Marks] &bull; {paper} ({year})\n\n"
            full_markdown += f"{q_latex}\n\n"
            if sol_latex:
                clean_sol = normalize_math_delimiters(sol_latex.strip())
                full_markdown += "> **✓ Step-by-Step Verified Solution & Marking Rubric:**\n>\n"
                sol_lines = clean_sol.split("\n")
                full_markdown += "\n".join(f"> {line}" for line in sol_lines) + "\n\n"

    if practice_questions:
        full_markdown += "\n\n---\n\n## Curated Practice Examination Variants & Solutions\n\n"
        for q in practice_questions:
            num = q.get("number", 1)
            marks = q.get("marks", 10)
            topic = q.get("topic", topic_title)
            q_latex = q.get("question_latex", "")
            sol_latex = q.get("solution_latex", "")
            full_markdown += f"### Practice Question {num} [{marks} Marks] &bull; {topic}\n\n"
            full_markdown += f"{q_latex}\n\n"
            if sol_latex:
                clean_sol = normalize_math_delimiters(sol_latex.strip())
                full_markdown += "> **✓ Step-by-Step Solution & Marking Scheme:**\n>\n"
                sol_lines = clean_sol.split("\n")
                full_markdown += "\n".join(f"> {line}" for line in sol_lines) + "\n\n"

    date_str = datetime.now().strftime("%B %d, %Y")
    course_upper = course_code.upper()

    full_markdown = normalize_math_delimiters(full_markdown)
    escaped_json = json.dumps(full_markdown).replace("</", "<\/")

    level_labels = {
        "level_1": "Level 1: Intuition & Foundations",
        "level_2": "Core Concepts",
        "level_3": "Level 3: Exam Mode & High Yield",
    }
    level_display = level_labels.get(level, "Core Concepts")

    template_path = os.path.join(settings.BASE_DIR, "templates", "prep", "pdf_export_template.html")
    try:
        with open(template_path, "r", encoding="utf-8") as f:
            template_html = f.read()
    except Exception as e:
        logger.error("Could not load pdf_export_template.html: %s", e)
        template_html = "<html><body><div id='content'></div></body></html>"

    # Use the same current Markdown and KaTeX renderer as the study page.
    # It is inserted after the template's legacy fallback renderer, so this
    # shared source is the final owner of PDF rendering.
    renderer_path = os.path.join(settings.BASE_DIR, "static", "js", "prep-renderer.js")
    try:
        with open(renderer_path, "r", encoding="utf-8") as f:
            renderer_script = f.read()
    except OSError as exc:
        logger.warning("Could not load shared Prep renderer for PDF export: %s", exc)
        renderer_script = ""

    return (
        template_html.replace("{{COURSE_UPPER}}", course_upper)
        .replace("{{TOPIC_TITLE}}", topic_title)
        .replace("{{LEVEL_DISPLAY}}", level_display)
        .replace("{{DATE_STR}}", date_str)
        .replace("{{ESCAPED_JSON}}", escaped_json)
        .replace("{{RENDERER_JS}}", renderer_script)
    )

def _build_paper_questions_html(
    course_code: str,
    paper_title: str,
    year: str,
    total_marks: int,
    questions: list[dict],
    include_questions: bool = True,
    include_answers: bool = True,
) -> str:
    """Build a question paper, answer key, or combined paper from one source."""
    from services.prep_ai_router import normalize_math_delimiters

    course_upper = course_code.upper()
    if include_questions and include_answers:
        document_label = paper_title
    elif include_questions:
        document_label = f"{paper_title} - Question Paper"
    else:
        document_label = f"{paper_title} - Answer Key"
    questions_markdown = f"# {document_label}\n\n**Academic Period:** {year} | **Total Marks:** {total_marks} Marks\n\n---\n\n"
    for q in questions:
        num = q.get("number", 1)
        marks = q.get("marks", 10)
        topic = q.get("topic", "Mathematical Assessment")
        q_latex = normalize_math_delimiters(q.get("question_latex", ""))
        sol_latex = normalize_math_delimiters(q.get("solution_latex", ""))

        heading = f"## Question {num} [{marks} Marks] - {topic}"
        if include_answers and not include_questions:
            heading = f"## Answer {num} [{marks} Marks] - {topic}"
        questions_markdown += f"{heading}\n\n"
        if include_questions:
            questions_markdown += f"{q_latex}\n\n"
        if include_answers and sol_latex:
            questions_markdown += "### Step-by-Step Verified Solution\n\n"
            questions_markdown += f"{sol_latex.strip()}\n\n"
        elif include_answers:
            questions_markdown += "Solution status: A verified solution is not available for this question.\n\n"
        questions_markdown += "---\n\n"

    return _build_topic_notes_html(
        course_code=course_upper,
        topic_title=document_label,
        notes_content=questions_markdown,
    )


# ─── 2. Fallback PDF Export Engine (ReportLab) ────────────────────────────────

def _export_topic_notes_reportlab(
    course_code: str,
    topic_title: str,
    notes_content: str,
    authentic_questions: list | None = None,
    practice_questions: list | None = None,
) -> bytes:
    """Generate a clean, styled PDF for topic notes using ReportLab (fallback)."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()
    primary_color = colors.HexColor("#0f766e")  # Mentify Teal
    danger_color = colors.HexColor("#d93025")   # Prep Red
    text_color = colors.HexColor("#1e293b")
    muted_color = colors.HexColor("#64748b")
    bg_surface = colors.HexColor("#f8fafc")
    bg_code = colors.HexColor("#1e293b")

    title_style = ParagraphStyle(
        "PrepTitle",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=18,
        leading=22,
        textColor=primary_color,
    )
    subtitle_style = ParagraphStyle(
        "PrepSubtitle",
        parent=styles["Normal"],
        fontName=FONT_BODY,
        fontSize=9.5,
        leading=13,
        textColor=muted_color,
    )
    h1_style = ParagraphStyle(
        "PrepH1",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=14,
        leading=18,
        textColor=primary_color,
        spaceBefore=16,
        spaceAfter=8,
    )
    h2_style = ParagraphStyle(
        "PrepH2",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=12,
        leading=16,
        textColor=danger_color,
        spaceBefore=14,
        spaceAfter=6,
    )
    h3_style = ParagraphStyle(
        "PrepH3",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=10.5,
        leading=14,
        textColor=text_color,
        spaceBefore=10,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "PrepBody",
        parent=styles["Normal"],
        fontName=FONT_BODY,
        fontSize=9.5,
        leading=14.5,
        textColor=text_color,
    )
    math_style = ParagraphStyle(
        "PrepMath",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=10,
        leading=15,
        textColor=colors.HexColor("#0f172a"),
        alignment=1,  # Center
    )
    code_style = ParagraphStyle(
        "PrepCode",
        parent=styles["Normal"],
        fontName=FONT_CODE,
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#f8fafc"),
    )
    bq_style = ParagraphStyle(
        "PrepBQ",
        parent=styles["Normal"],
        fontName=FONT_ITALIC,
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor("#0f172a"),
    )
    table_cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName=FONT_BODY,
        fontSize=8.5,
        leading=11.5,
        textColor=text_color,
    )
    table_header_style = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=8.5,
        leading=11.5,
        textColor=colors.white,
    )
    q_title_style = ParagraphStyle(
        "QTitle",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=10.5,
        leading=14,
        textColor=danger_color,
        spaceBefore=12,
        spaceAfter=4,
    )
    sol_header_style = ParagraphStyle(
        "SolHeader",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#16a34a"),
    )

    story = []

    # Header Banner
    story.append(Paragraph(f"MENTIFY PREP &bull; {course_code.upper()}", subtitle_style))
    story.append(Spacer(1, 3))
    story.append(Paragraph(topic_title, title_style))
    story.append(Paragraph(f"Canonical Syllabus Notes & Exam Preparation &bull; Generated {datetime.now().strftime('%B %d, %Y')}", subtitle_style))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1.5, color=primary_color, spaceBefore=2, spaceAfter=12))

    # Parse and build blocks
    blocks = _parse_markdown_into_blocks(notes_content)

    for block in blocks:
        b_type = block["type"]

        if b_type == "heading":
            level = block["level"]
            text = _format_markdown_for_reportlab(block["text"])
            if level == 1:
                story.append(Paragraph(text, h1_style))
            elif level == 2:
                story.append(Paragraph(text, h2_style))
            else:
                story.append(Paragraph(text, h3_style))

        elif b_type == "paragraph":
            story.append(Paragraph(_format_markdown_for_reportlab(block["text"]), body_style))
            story.append(Spacer(1, 4))

        elif b_type == "math":
            beautified = beautify_math_formula(block["formula"])
            # Format XML-safe
            b_xml = beautified.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            p_math = Paragraph(b_xml, math_style)
            t_math = Table([[p_math]], colWidths=["100%"])
            t_math.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), bg_surface),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ]))
            story.append(Spacer(1, 4))
            story.append(t_math)
            story.append(Spacer(1, 6))

        elif b_type == "code":
            raw_code = "<br/>".join(
                line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace(" ", "&nbsp;")
                for line in block["lines"]
            )
            p_code = Paragraph(raw_code, code_style)
            t_code = Table([[p_code]], colWidths=["100%"])
            t_code.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), bg_code),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#0f172a")),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ]))
            story.append(Spacer(1, 4))
            story.append(t_code)
            story.append(Spacer(1, 6))

        elif b_type == "table":
            headers = block["headers"]
            rows = block["rows"]
            num_cols = len(headers)
            if num_cols > 0:
                col_width = (doc.width) / num_cols
                col_widths = [col_width] * num_cols

                table_data = []
                # Header row
                table_data.append([Paragraph(_format_markdown_for_reportlab(h), table_header_style) for h in headers])
                # Data rows
                for r in rows:
                    # Pad row if missing columns
                    padded = r + [""] * (num_cols - len(r))
                    table_data.append([Paragraph(_format_markdown_for_reportlab(c), table_cell_style) for c in padded[:num_cols]])

                tbl = Table(table_data, colWidths=col_widths)
                tbl.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), primary_color),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, bg_surface]),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ]))
                story.append(Spacer(1, 4))
                story.append(tbl)
                story.append(Spacer(1, 8))

        elif b_type == "blockquote":
            p_bq = Paragraph(_format_markdown_for_reportlab(block["text"]), bq_style)
            t_bq = Table([[p_bq]], colWidths=["100%"])
            t_bq.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), bg_surface),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                ("LINEBEFORE", (0, 0), (0, -1), 3, primary_color),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ]))
            story.append(Spacer(1, 4))
            story.append(t_bq)
            story.append(Spacer(1, 6))

    # Append Authentic Questions & Verified Proofs
    if authentic_questions:
        story.append(Spacer(1, 14))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceBefore=6, spaceAfter=12))
        story.append(Paragraph("Authentic Examination Problems &amp; Verified Proofs", h1_style))
        for q in authentic_questions:
            num = q.get("number", 1)
            marks = q.get("marks", 10)
            paper = q.get("paper_title", "University Examination")
            year = q.get("year", "")
            q_latex = q.get("question_latex", "")
            sol_latex = q.get("solution_latex", "")

            story.append(Paragraph(f"Question {num} [{marks} Marks] &bull; {paper} ({year})", q_title_style))
            story.append(Paragraph(_format_markdown_for_reportlab(q_latex), body_style))
            if sol_latex:
                sol_blocks = _parse_markdown_into_blocks(sol_latex)
                sol_flowables = [Paragraph("<b>Step-by-Step Verified Solution &amp; Marking Rubric:</b>", sol_header_style)]
                for sb in sol_blocks:
                    if sb["type"] == "code":
                        c_text = "<br/>".join(
                            l.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace(" ", "&nbsp;")
                            for l in sb["lines"]
                        )
                        sol_flowables.append(Paragraph(c_text, code_style))
                    elif sb["type"] == "math":
                        sol_flowables.append(Paragraph(beautify_math_formula(sb["formula"]), math_style))
                    else:
                        sol_flowables.append(Paragraph(_format_markdown_for_reportlab(sb.get("text", "")), body_style))

                sol_table = Table([[f] for f in sol_flowables], colWidths=["100%"])
                sol_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, -1), bg_surface),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ]))
                story.append(Spacer(1, 4))
                story.append(sol_table)
            story.append(Spacer(1, 10))

    # Append Practice Variants if present
    if practice_questions:
        story.append(Spacer(1, 14))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceBefore=6, spaceAfter=12))
        story.append(Paragraph("Curated Practice Examination Variants &amp; Solutions", h1_style))
        for q in practice_questions:
            num = q.get("number", 1)
            marks = q.get("marks", 10)
            topic = q.get("topic", topic_title)
            q_latex = q.get("question_latex", "")
            sol_latex = q.get("solution_latex", "")

            story.append(Paragraph(f"Practice Question {num} [{marks} Marks] &bull; {topic}", q_title_style))
            story.append(Paragraph(_format_markdown_for_reportlab(q_latex), body_style))
            if sol_latex:
                sol_table = Table([[
                    Paragraph("<b>Step-by-Step Solution &amp; Marking Scheme:</b>", sol_header_style)
                ], [
                    Paragraph(_format_markdown_for_reportlab(sol_latex).replace("\n", "<br/>"), body_style)
                ]], colWidths=["100%"])
                sol_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, -1), bg_surface),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ]))
                story.append(Spacer(1, 4))
                story.append(sol_table)
            story.append(Spacer(1, 10))

    doc.build(story)
    return buffer.getvalue()


def _export_paper_questions_reportlab(
    course_code: str,
    paper_title: str,
    year: str,
    total_marks: int,
    questions: list[dict],
    include_questions: bool = True,
    include_answers: bool = True,
) -> bytes:
    """Generate a clean, styled PDF for CAT examination questions using ReportLab (fallback)."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()
    primary_color = colors.HexColor("#0f766e")
    danger_color = colors.HexColor("#d93025")
    text_color = colors.HexColor("#1e293b")
    muted_color = colors.HexColor("#64748b")
    bg_surface = colors.HexColor("#f8fafc")

    title_style = ParagraphStyle(
        "PaperTitle",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=17,
        leading=21,
        textColor=danger_color,
    )
    meta_style = ParagraphStyle(
        "PaperMeta",
        parent=styles["Normal"],
        fontName=FONT_BODY,
        fontSize=9,
        leading=13,
        textColor=muted_color,
    )
    q_header_style = ParagraphStyle(
        "QHeader",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=10.5,
        leading=14,
        textColor=primary_color,
    )
    q_marks_style = ParagraphStyle(
        "QMarks",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=9,
        leading=12,
        textColor=danger_color,
        alignment=2,
    )
    body_style = ParagraphStyle(
        "QBody",
        parent=styles["Normal"],
        fontName=FONT_BODY,
        fontSize=9.5,
        leading=14,
        textColor=text_color,
    )
    proof_header_style = ParagraphStyle(
        "ProofHeader",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#16a34a"),
    )
    code_style = ParagraphStyle(
        "ProofCode",
        parent=styles["Normal"],
        fontName=FONT_CODE,
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#f8fafc"),
    )

    story = []

    # Paper Top Header
    story.append(Paragraph(f"MENTIFY PREP &bull; {course_code.upper()}", meta_style))
    story.append(Spacer(1, 3))
    document_label = paper_title
    if include_questions and not include_answers:
        document_label = f"{paper_title} - Question Paper"
    elif include_answers and not include_questions:
        document_label = f"{paper_title} - Answer Key"
    story.append(Paragraph(document_label, title_style))
    story.append(Paragraph(f"Academic Period: {year} &bull; Total Marks: {total_marks} MARKS", meta_style))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1.5, color=danger_color, spaceBefore=2, spaceAfter=12))

    # Questions loop
    for q in questions:
        num = q.get("number", 1)
        marks = q.get("marks", 10)
        topic = q.get("topic", "")
        question_latex = q.get("question_latex", "")
        solution_latex = q.get("solution_latex", "")

        q_flowables = []

        header_table = Table(
            [[
                Paragraph(f"QUESTION {num}: {topic}", q_header_style),
                Paragraph(f"[{marks} MARKS]", q_marks_style),
            ]],
            colWidths=["80%", "20%"],
        )
        header_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        q_flowables.append(header_table)
        q_flowables.append(Spacer(1, 3))

        # Omit problem statements from answer-only exports.
        if include_questions:
            q_flowables.append(Paragraph(_format_markdown_for_reportlab(question_latex), body_style))
            q_flowables.append(Spacer(1, 6))

        # Verified Proof / Solution
        if include_answers and solution_latex:
            sol_blocks = _parse_markdown_into_blocks(solution_latex)
            sol_flowables = [Paragraph("&check; VERIFIED STEP-BY-STEP PROOF / SOLUTION:", proof_header_style)]
            for sb in sol_blocks:
                if sb["type"] == "code":
                    c_text = "<br/>".join(
                        l.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace(" ", "&nbsp;")
                        for l in sb["lines"]
                    )
                    sol_flowables.append(Paragraph(c_text, code_style))
                elif sb["type"] == "math":
                    sol_flowables.append(Paragraph(beautify_math_formula(sb["formula"]), body_style))
                else:
                    sol_flowables.append(Paragraph(_format_markdown_for_reportlab(sb.get("text", "")), body_style))

            t_sol = Table([[f] for f in sol_flowables], colWidths=["100%"])
            t_sol.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), bg_surface),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ]))
            q_flowables.append(t_sol)
            q_flowables.append(Spacer(1, 8))
        elif include_answers:
            q_flowables.append(Paragraph("No verified solution is available for this question.", body_style))
            q_flowables.append(Spacer(1, 8))

        q_flowables.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#e2e8f0"), spaceBefore=4, spaceAfter=12))
        story.append(KeepTogether(q_flowables))

    doc.build(story)
    return buffer.getvalue()


# ─── Public PDF Export Interface ─────────────────────────────────────────────

def export_topic_notes_pdf(
    course_code: str,
    topic_title: str,
    notes_content: str,
    authentic_questions: list | None = None,
    practice_questions: list | None = None,
    level: str = "level_2",
) -> bytes:
    """
    Generate publication-ready PDF for topic notes.
    Uses Playwright + Chrome for pixel-perfect KaTeX vectors and syntax highlighting,
    with automatic fallback to ReportLab.
    """
    level_labels = {
        "level_1": "Level 1: Intuition & Foundations",
        "level_2": "Core Concepts",
        "level_3": "Level 3: Exam Mode & High Yield",
    }
    level_display = level_labels.get(level, "Core Concepts")

    try:
        html = _build_topic_notes_html(
            course_code=course_code,
            topic_title=topic_title,
            notes_content=notes_content,
            authentic_questions=authentic_questions,
            practice_questions=practice_questions,
            level=level,
        )
        pdf_bytes = _render_html_with_playwright(
            html_content=html,
            header_text=f"{course_code.upper()} • {topic_title} • {level_display}",
        )
        if pdf_bytes and len(pdf_bytes) > 1000:
            return pdf_bytes
    except Exception as exc:
        logger.warning("Playwright topic PDF export failed, falling back to ReportLab: %s", exc)

    return _export_topic_notes_reportlab(
        course_code=course_code,
        topic_title=topic_title,
        notes_content=notes_content,
        authentic_questions=authentic_questions,
        practice_questions=practice_questions,
    )


def export_paper_questions_pdf(
    course_code: str,
    paper_title: str,
    year: str,
    total_marks: int,
    questions: list[dict],
    include_questions: bool = True,
    include_answers: bool = True,
) -> bytes:
    """
    Generate publication-ready PDF for past paper questions and step-by-step proofs.
    Uses Playwright + Chrome for pixel-perfect KaTeX vectors, with automatic fallback to ReportLab.
    """
    if not include_questions and not include_answers:
        raise ValueError("At least questions or answers must be selected for export.")

    try:
        html = _build_paper_questions_html(
            course_code=course_code,
            paper_title=paper_title,
            year=year,
            total_marks=total_marks,
            questions=questions,
            include_questions=include_questions,
            include_answers=include_answers,
        )
        pdf_bytes = _render_html_with_playwright(
            html_content=html,
            header_text=f"{course_code.upper()} • {paper_title}",
        )
        if pdf_bytes and len(pdf_bytes) > 1000:
            return pdf_bytes
    except Exception as exc:
        logger.warning("Playwright paper PDF export failed, falling back to ReportLab: %s", exc)

    return _export_paper_questions_reportlab(
        course_code=course_code,
        paper_title=paper_title,
        year=year,
        total_marks=total_marks,
        questions=questions,
        include_questions=include_questions,
        include_answers=include_answers,
    )


# ─── 3. DOCX Export Engine (python-docx) ──────────────────────────────────────

def _set_cell_background(cell, hex_color: str):
    """Set background color of a Word table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _style_callout_cell(cell, border_color="0F766E", bg_color="F8FAFC"):
    """Style Word table cell as a callout with left accent border and soft background."""
    _set_cell_background(cell, bg_color)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)


def _style_code_cell(cell):
    """Style Word table cell as a code block with subtle border and dark slate background."""
    _set_cell_background(cell, "1E293B")
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="334155"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="334155"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="334155"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="334155"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)


def _add_styled_runs_to_paragraph(p, text: str):
    """Adds styled runs (bold, italic, code, beautified math) to a Word paragraph."""
    pattern = r"(\$[^\$\n\r]+?\$|\*\*[^*]+?\*\*|`[^`]+?`|\*[^*]+?\*)"
    parts = re.split(pattern, text)
    for part in parts:
        if not part:
            continue
        if part.startswith("$") and part.endswith("$") and len(part) > 1:
            r = p.add_run(beautify_math_formula(part))
            r.font.name = "Cambria Math"
            r.font.bold = True
            r.font.italic = True
        elif part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.font.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(15, 118, 110)
        elif part.startswith("*") and part.endswith("*"):
            r = p.add_run(part[1:-1])
            r.font.italic = True
        else:
            p.add_run(part)


def _add_styled_paragraph(doc, text: str, style_name="Normal", space_after=4):
    """Adds a paragraph with inline bold/italic and beautified math to Word document."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(space_after)
    _add_styled_runs_to_paragraph(p, text)
    return p


def export_topic_notes_docx(
    course_code: str,
    topic_title: str,
    notes_content: str,
    authentic_questions: list | None = None,
    practice_questions: list | None = None,
) -> bytes:
    """Generate a clean Microsoft Word DOCX document with formatted math, tables, code, and exam variants."""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run(f"MENTIFY PREP • {course_code.upper()} COMPREHENSIVE STUDY PACK")
    r_meta.font.size = Pt(9)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    h_main = doc.add_heading(topic_title, level=1)
    for r in h_main.runs:
        r.font.color.rgb = RGBColor(15, 118, 110)

    p_sub = doc.add_paragraph(f"Canonical Syllabus Notes • Generated {datetime.now().strftime('%B %d, %Y')}")
    p_sub.runs[0].font.size = Pt(9.5)
    p_sub.runs[0].font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Parse markdown into structured blocks
    blocks = _parse_markdown_into_blocks(notes_content)

    for b in blocks:
        b_type = b["type"]

        if b_type == "heading":
            level = b["level"]
            text = b["text"]
            h = doc.add_heading(text, level=min(level, 3))
            color = RGBColor(15, 118, 110) if level <= 2 else RGBColor(217, 48, 37)
            for r in h.runs:
                r.font.color.rgb = color

        elif b_type == "paragraph":
            _add_styled_paragraph(doc, b["text"])

        elif b_type == "math":
            formula = beautify_math_formula(b["formula"])
            p_math = doc.add_paragraph()
            p_math.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_math.paragraph_format.space_before = Pt(6)
            p_math.paragraph_format.space_after = Pt(6)
            p_math.paragraph_format.left_indent = Inches(0.4)
            p_math.paragraph_format.right_indent = Inches(0.4)
            r_math = p_math.add_run(formula)
            r_math.font.name = "Cambria Math"
            r_math.font.size = Pt(11)
            r_math.font.bold = True
            r_math.font.color.rgb = RGBColor(15, 23, 42)

        elif b_type == "code":
            tbl = doc.add_table(rows=1, cols=1)
            cell = tbl.cell(0, 0)
            _style_code_cell(cell)
            p_cell = cell.paragraphs[0]
            p_cell.paragraph_format.line_spacing = 1.0
            p_cell.paragraph_format.space_before = Pt(4)
            p_cell.paragraph_format.space_after = Pt(4)
            code_text = "\n".join(b["lines"])
            r_code = p_cell.add_run(code_text)
            r_code.font.name = "Consolas"
            r_code.font.size = Pt(9)
            r_code.font.color.rgb = RGBColor(248, 250, 252)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        elif b_type == "table":
            headers = b["headers"]
            rows = b["rows"]
            num_cols = len(headers)
            if num_cols > 0:
                t = doc.add_table(rows=len(rows) + 1, cols=num_cols)
                t.style = "Table Grid"
                # Header row
                for idx, h_text in enumerate(headers):
                    cell = t.cell(0, idx)
                    _set_cell_background(cell, "0F766E")
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    r = p.add_run(h_text)
                    r.font.bold = True
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(255, 255, 255)
                # Data rows
                for r_idx, r_data in enumerate(rows, start=1):
                    for c_idx in range(num_cols):
                        val = r_data[c_idx] if c_idx < len(r_data) else ""
                        cell = t.cell(r_idx, c_idx)
                        if r_idx % 2 == 0:
                            _set_cell_background(cell, "F8FAFC")
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        _add_styled_runs_to_paragraph(p, val)
                doc.add_paragraph().paragraph_format.space_after = Pt(6)

        elif b_type == "blockquote":
            tbl = doc.add_table(rows=1, cols=1)
            cell = tbl.cell(0, 0)
            _style_callout_cell(cell, "0F766E")
            p_cell = cell.paragraphs[0]
            p_cell.paragraph_format.space_before = Pt(4)
            p_cell.paragraph_format.space_after = Pt(4)
            _add_styled_runs_to_paragraph(p_cell, b["text"])
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Append Authentic Questions
    if authentic_questions:
        doc.add_paragraph().paragraph_format.space_after = Pt(12)
        h_auth = doc.add_heading("Authentic Examination Problems & Verified Proofs", level=1)
        for r in h_auth.runs:
            r.font.color.rgb = RGBColor(15, 118, 110)

        for q in authentic_questions:
            num = q.get("number", 1)
            marks = q.get("marks", 10)
            paper = q.get("paper_title", "University Examination")
            year = q.get("year", "")
            q_latex = q.get("question_latex", "")
            sol_latex = q.get("solution_latex", "")

            p_qh = doc.add_paragraph()
            r_qh = p_qh.add_run(f"Question {num} [{marks} Marks] • {paper} ({year})")
            r_qh.font.bold = True
            r_qh.font.size = Pt(10.5)
            r_qh.font.color.rgb = RGBColor(217, 48, 37)

            _add_styled_paragraph(doc, q_latex)

            if sol_latex:
                sol_table = doc.add_table(rows=1, cols=1)
                sol_cell = sol_table.cell(0, 0)
                _style_callout_cell(sol_cell, "16A34A")
                p_head = sol_cell.paragraphs[0]
                r_sol_h = p_head.add_run("✓ Step-by-Step Verified Solution & Marking Rubric:")
                r_sol_h.font.bold = True
                r_sol_h.font.size = Pt(9.5)
                r_sol_h.font.color.rgb = RGBColor(22, 163, 74)

                p_body = sol_cell.add_paragraph()
                _add_styled_runs_to_paragraph(p_body, beautify_math_formula(sol_latex))
                p_body.paragraph_format.line_spacing = 1.15

            doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Append Practice Variants
    if practice_questions:
        doc.add_paragraph().paragraph_format.space_after = Pt(12)
        h_prac = doc.add_heading("Curated Practice Examination Variants & Solutions", level=1)
        for r in h_prac.runs:
            r.font.color.rgb = RGBColor(15, 118, 110)

        for q in practice_questions:
            num = q.get("number", 1)
            marks = q.get("marks", 10)
            topic = q.get("topic", topic_title)
            q_latex = q.get("question_latex", "")
            sol_latex = q.get("solution_latex", "")

            p_qh = doc.add_paragraph()
            r_qh = p_qh.add_run(f"Practice Question {num} [{marks} Marks] • {topic}")
            r_qh.font.bold = True
            r_qh.font.size = Pt(10.5)
            r_qh.font.color.rgb = RGBColor(217, 48, 37)

            _add_styled_paragraph(doc, q_latex)

            if sol_latex:
                sol_table = doc.add_table(rows=1, cols=1)
                sol_cell = sol_table.cell(0, 0)
                _style_callout_cell(sol_cell, "16A34A")
                p_head = sol_cell.paragraphs[0]
                r_sol_h = p_head.add_run("✓ Step-by-Step Solution & Marking Scheme:")
                r_sol_h.font.bold = True
                r_sol_h.font.size = Pt(9.5)
                r_sol_h.font.color.rgb = RGBColor(22, 163, 74)

                p_body = sol_cell.add_paragraph()
                _add_styled_runs_to_paragraph(p_body, beautify_math_formula(sol_latex))
                p_body.paragraph_format.line_spacing = 1.15

            doc.add_paragraph().paragraph_format.space_after = Pt(6)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def export_paper_questions_docx(
    course_code: str,
    paper_title: str,
    year: str,
    total_marks: int,
    questions: list[dict],
) -> bytes:
    """Generate a clean Microsoft Word DOCX document for CAT papers with solution proofs."""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Header
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run(f"MENTIFY PREP • {course_code.upper()} ASSESSMENT ARCHIVE")
    r_meta.font.size = Pt(9)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    h_main = doc.add_heading(paper_title, level=1)
    for r in h_main.runs:
        r.font.color.rgb = RGBColor(217, 48, 37)

    p_sub = doc.add_paragraph(f"Academic Period: {year} | Total Marks: {total_marks} Marks")
    p_sub.runs[0].font.size = Pt(9.5)
    p_sub.runs[0].font.bold = True

    for q in questions:
        num = q.get("number", 1)
        marks = q.get("marks", 10)
        topic = q.get("topic", "")
        question_latex = q.get("question_latex", "")
        solution_latex = q.get("solution_latex", "")

        q_table = doc.add_table(rows=1, cols=2)
        cell_l = q_table.cell(0, 0)
        cell_r = q_table.cell(0, 1)
        _set_cell_background(cell_l, "F1F5F9")
        _set_cell_background(cell_r, "F1F5F9")

        p_l = cell_l.paragraphs[0]
        r_l = p_l.add_run(f"QUESTION {num}: {topic}")
        r_l.font.bold = True
        r_l.font.size = Pt(10.5)
        r_l.font.color.rgb = RGBColor(15, 118, 110)

        p_r = cell_r.paragraphs[0]
        p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_r = p_r.add_run(f"[{marks} MARKS]")
        r_r.font.bold = True
        r_r.font.size = Pt(10)
        r_r.font.color.rgb = RGBColor(217, 48, 37)

        _add_styled_paragraph(doc, question_latex, space_after=6)

        if solution_latex:
            sol_table = doc.add_table(rows=1, cols=1)
            sol_cell = sol_table.cell(0, 0)
            _style_callout_cell(sol_cell, "16A34A")
            p_head = sol_cell.paragraphs[0]
            r_sol_h = p_head.add_run("✓ Verified Step-by-Step Proof / Solution:")
            r_sol_h.font.bold = True
            r_sol_h.font.size = Pt(9.5)
            r_sol_h.font.color.rgb = RGBColor(22, 163, 74)

            p_body = sol_cell.add_paragraph()
            _add_styled_runs_to_paragraph(p_body, beautify_math_formula(solution_latex))
            p_body.paragraph_format.line_spacing = 1.15

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
